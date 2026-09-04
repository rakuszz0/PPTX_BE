from __future__ import annotations
from typing import TYPE_CHECKING
import os

from pptx import Presentation as PptxPresentation

from app.domain.presentation.models import PresentationDocument, Slide, Component, ComponentType
from app.services.design.theme import Theme, ThemeSystem
from .pptx_utils import (
    add_rectangle, add_textbox, set_textbox, set_multibullet,
    set_slide_background, add_shape_band, set_speaker_notes,
    rgb_from_hex,
)

if TYPE_CHECKING:
    from pptx.parts.slide import SlidePart


class PptxRenderer:
    def __init__(self, theme: Theme | None = None):
        self.theme = theme or ThemeSystem.get("medical_professional")

    def _c(self, name: str, fallback: str = "#000000") -> str:
        return self.theme.get_color_hex(name, fallback)

    def render(self, document: PresentationDocument, output_path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)

        prs = PptxPresentation()
        prs.slide_width = Inches_round(document.settings.slide_width_in)
        prs.slide_height = Inches_round(document.settings.slide_height_in)

        blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]

        total = len(document.slides)
        for idx, slide in enumerate(document.slides):
            ps = prs.slides.add_slide(blank_layout)
            self._render_slide(ps, slide, idx + 1, total, document)

        prs.save(output_path)
        return output_path

    def _render_slide(self, ps: "SlidePart", slide: Slide, index: int, total: int,
                      doc: PresentationDocument) -> None:
        set_slide_background(ps, self._c("background", "#FBFCFE"))
        self._add_slide_chrome(ps, index, total, doc)

        for comp in slide.components:
            self._render_component(ps, comp)

        set_speaker_notes(ps, slide.speaker_notes)

    def _add_slide_chrome(self, ps: "SlidePart", index: int, total: int, doc: PresentationDocument) -> None:
        from app.domain.presentation.models import LayoutType
        lt = slide_layout_type(slide_layout_from_str(None))
        title_like = False
        try:
            _ = LayoutType.TITLE
        except Exception:
            pass
        if hasattr(ps, "_wizape_layout"):
            title_like = ps._wizape_layout in ("TITLE", "SECTION", "BIG_STATEMENT")
        if index == 1:
            title_like = True

        if not title_like:
            accent_h = 0.08
            add_shape_band(ps, 0.0, accent_h, self._c("primary", "#165DFF"), width_in=doc.settings.slide_width_in)
            footer_y = doc.settings.slide_height_in - 0.35
            add_rectangle(ps, 0, footer_y, doc.settings.slide_width_in, 0.35,
                          fill_hex=self._c("neutral_50", "#F8F9FC"), line_hex=None, radius=0.0)
            if doc.settings.footer:
                tb = add_textbox(ps, 0.4, footer_y + 0.07, 8.0, 0.22)
                set_textbox(tb, doc.settings.footer, size_pt=10,
                            color_hex=self._c("footer", "#36435C"), italic=True)
            if doc.settings.show_page_numbers:
                tb = add_textbox(ps, doc.settings.slide_width_in - 1.6, footer_y + 0.07, 1.2, 0.22)
                set_textbox(tb, f"{index} / {total}", size_pt=10,
                            color_hex=self._c("page_number", "#4E5A75"), align="right")

    def _render_component(self, ps: "SlidePart", comp: Component) -> None:
        x, y, w, h = comp.position.x, comp.position.y, comp.size.width, comp.size.height
        if w <= 0 or h <= 0:
            return
        c = comp.content
        st = comp.style
        ty = comp.type
        bg = (st.background and st.background.to_hex()) or None
        bd = (st.border_color and st.border_color.to_hex()) or None
        bw = st.border_width or 0.75
        br = st.border_radius or 0.06

        if ty in (ComponentType.HEADING, ComponentType.SUBTITLE):
            tb = add_textbox(ps, x, y, w, h)
            preset = "title" if ty == ComponentType.HEADING and (h > 0.8 or w > 9) else (
                "heading_1" if ty == ComponentType.HEADING else "subtitle"
            )
            typ = self.theme.typography_system.get(preset)
            set_textbox(tb, c.text or c.heading or "",
                        size_pt=typ.size_pt if typ else (32 if ty == ComponentType.HEADING else 20),
                        bold=typ.weight in ("Bold", "SemiBold") if typ else True,
                        color_hex=(st.text and st.text.color and st.text.color.to_hex()) or self._c(typ.color_name if typ else "text_heading"),
                        font_family=typ.font_family if typ else "Calibri",
                        align=(st.text and st.text.align) or "left",
                        line_spacing=typ.line_height if typ else 1.2)
            return

        if ty == ComponentType.PARAGRAPH:
            tb = add_textbox(ps, x, y, w, h)
            typ = self.theme.typography_system.get("body")
            set_textbox(tb, c.text or "",
                        size_pt=typ.size_pt if typ else 14,
                        color_hex=self._c(typ.color_name if typ else "text_body"),
                        font_family=typ.font_family if typ else "Calibri",
                        line_spacing=typ.line_height if typ else 1.45,
                        align=(st.text and st.text.align) or "left")
            return

        if ty == ComponentType.BULLET_LIST:
            if bg:
                add_rectangle(ps, x, y, w, h, fill_hex=bg, line_hex=bd, line_width_pt=bw, radius=br)
                x += 0.12; y += 0.08; w -= 0.24; h -= 0.16
            tb = add_textbox(ps, x, y, w, h)
            typ = self.theme.typography_system.get("bullet")
            if c.items:
                set_multibullet(tb, c.items,
                                size_pt=typ.size_pt if typ else 14,
                                color_hex=self._c(typ.color_name if typ else "text_body"),
                                font_family=typ.font_family if typ else "Calibri",
                                bullet_color_hex=self._c("primary", "#165DFF"))
            return

        if ty == ComponentType.CARD:
            variant = st.variant or "card"
            ss = self.theme.shape_system.get(variant)
            fill_hex = bg or self._c(ss.fill if ss else "surface")
            line_hex = bd or (self._c(ss.border) if ss else None)
            shape = add_rectangle(ps, x, y, w, h,
                                  fill_hex=fill_hex, line_hex=line_hex,
                                  line_width_pt=bw or (ss.border_width_pt if ss else 0.75),
                                  radius=br or (ss.border_radius_in if ss else 0.1))
            if st.elevation or (ss and ss.shadow):
                try:
                    shape.shadow.inherit = True
                    shape.shadow.blur_radius = 1
                except Exception:
                    pass
            title_h = min(h * 0.28, 0.55)
            if c.heading:
                tb = add_textbox(ps, x + 0.22, y + 0.18, w - 0.44, title_h)
                typ = self.theme.typography_system.get("heading_3")
                set_textbox(tb, c.heading,
                            size_pt=typ.size_pt if typ else 18, bold=True,
                            color_hex=self._c(typ.color_name if typ else "text_heading"),
                            align=(st.text and st.text.align) or "left")
            if c.text:
                by = y + 0.18 + title_h + 0.08
                bh = y + h - by - 0.18
                if bh > 0.25:
                    tb = add_textbox(ps, x + 0.22, by, w - 0.44, bh)
                    typ = self.theme.typography_system.get("body_small")
                    set_textbox(tb, c.text,
                                size_pt=typ.size_pt if typ else 12,
                                color_hex=self._c(typ.color_name if typ else "text_body"),
                                line_spacing=1.4)
            return

        if ty == ComponentType.METRIC:
            ss = self.theme.shape_system.get("card_primary")
            add_rectangle(ps, x, y, w, h,
                          fill_hex=bg or self._c(ss.fill if ss else "primary_light"),
                          line_hex=bd or self._c(ss.border if ss else "primary"),
                          line_width_pt=1.0, radius=br or 0.12)
            ty_v = self.theme.typography_system.get("metric_value")
            tb = add_textbox(ps, x + 0.18, y + 0.22, w - 0.36, h * 0.5)
            set_textbox(tb, str(c.value or ""),
                        size_pt=ty_v.size_pt if ty_v else 32, bold=True,
                        color_hex=(st.text and st.text.color and st.text.color.to_hex()) or self._c(ty_v.color_name if ty_v else "primary"),
                        align="center")
            if c.label:
                tb = add_textbox(ps, x + 0.18, y + h * 0.65, w - 0.36, h * 0.3)
                ty_l = self.theme.typography_system.get("metric_label")
                set_textbox(tb, c.label,
                            size_pt=ty_l.size_pt if ty_l else 12,
                            color_hex=self._c(ty_l.color_name if ty_l else "text_muted"),
                            align="center")
            return

        if ty == ComponentType.BADGE:
            variant = st.variant or "pill"
            ss = self.theme.shape_system.get(variant)
            fill_hex = bg or self._c(ss.fill if ss else "primary")
            line_hex = bd or (self._c(ss.border) if ss else None)
            add_rectangle(ps, x, y, w, h,
                          fill_hex=fill_hex, line_hex=line_hex,
                          line_width_pt=0, radius=br or (ss.border_radius_in if ss else 0.3))
            tb = add_textbox(ps, x + 0.08, y + 0.04, w - 0.16, h - 0.08)
            ty_b = self.theme.typography_system.get("badge")
            set_textbox(tb, c.text or "",
                        size_pt=ty_b.size_pt if ty_b else 10, bold=True,
                        color_hex=self._c("white", "#FFFFFF"), align="center")
            return

        if ty == ComponentType.CALLOUT:
            ss = self.theme.shape_system.get("callout")
            add_rectangle(ps, x, y, w, h,
                          fill_hex=bg or self._c(ss.fill if ss else "accent_teal_light"),
                          line_hex=bd or self._c(ss.border if ss else "accent_teal"),
                          line_width_pt=bw or 1.0,
                          radius=br or 0.08)
            tb = add_textbox(ps, x + 0.25, y + 0.18, w - 0.5, h - 0.36)
            set_textbox(tb, c.text or c.description or "",
                        size_pt=14, bold=False,
                        color_hex=self._c("text_body", "#1F2A44"), line_spacing=1.4)
            return

        if ty == ComponentType.QUOTE:
            mark_tb = add_textbox(ps, x, y, w * 0.12, h * 0.4)
            set_textbox(mark_tb, "\u201C", size_pt=72, bold=True,
                        color_hex=self._c("primary_light", "#E8F0FF"))
            tb = add_textbox(ps, x + w * 0.10, y + h * 0.1, w * 0.85, h * 0.7)
            ty_q = self.theme.typography_system.get("quote")
            set_textbox(tb, c.text or "",
                        size_pt=ty_q.size_pt if ty_q else 18,
                        italic=True,
                        color_hex=self._c(ty_q.color_name if ty_q else "neutral_700"),
                        font_family=ty_q.font_family if ty_q else "Georgia",
                        line_spacing=1.5)
            if c.author or c.citation:
                src = c.author or ""
                if c.citation:
                    src += f" \u2014 {c.citation}" if src else c.citation
                tb = add_textbox(ps, x + w * 0.10, y + h * 0.82, w * 0.85, h * 0.15)
                set_textbox(tb, src, size_pt=11, bold=True,
                            color_hex=self._c("neutral_600", "#4E5A75"))
            return

        if ty == ComponentType.TABLE:
            rows = c.rows or []
            headers = c.headers or []
            n_rows = len(rows) + (1 if headers else 0)
            n_cols = max([len(headers)] + [len(r) for r in rows] + [0])
            if n_rows == 0 or n_cols == 0:
                return
            table_shape = ps.shapes.add_table(n_rows, n_cols, Inches_round(x), Inches_round(y),
                                              Inches_round(w), Inches_round(h)).table
            if headers:
                for j, h in enumerate(headers[:n_cols]):
                    cell = table_shape.cell(0, j)
                    cell.text = str(h)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.bold = True
                            r.font.size = Pt(12)
                            r.font.color.rgb = rgb_from_hex(self._c("white", "#FFFFFF"))
                    from pptx.oxml.ns import qn
                    try:
                        tcPr = cell._tc.get_or_add_tcPr()
                        solidFill = tcPr.makeelement(qn('a:solidFill'), {})
                        srgbClr = solidFill.makeelement(qn('a:srgbClr'),
                                                        {'val': self._c("primary", "#165DFF").lstrip('#')})
                        solidFill.append(srgbClr)
                        tcPr.append(solidFill)
                    except Exception:
                        pass
            for i, row in enumerate(rows):
                ri = i + (1 if headers else 0)
                for j, v in enumerate(row[:n_cols]):
                    cell = table_shape.cell(ri, j)
                    cell.text = str(v)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(11)
                            r.font.color.rgb = rgb_from_hex(self._c("text_body", "#1F2A44"))
            return

        if ty in (ComponentType.FOOTER, ComponentType.PAGE_NUMBER, ComponentType.SOURCE_REFERENCE):
            tb = add_textbox(ps, x, y, w, h)
            ty_f = self.theme.typography_system.get("footer")
            set_textbox(tb, c.text or "",
                        size_pt=ty_f.size_pt if ty_f else 10,
                        color_hex=self._c(ty_f.color_name if ty_f else "footer"),
                        align=(st.text and st.text.align) or "left")
            return

        tb = add_textbox(ps, x, y, w, h)
        set_textbox(tb, c.text or "", size_pt=12,
                    color_hex=self._c("text_body", "#1F2A44"))


def Inches_round(n: float):
    from pptx.util import Inches
    return Inches(n)


def slide_layout_from_str(s):
    return s or "ONE_COLUMN"


def slide_layout_type(s):
    return s
