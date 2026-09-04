from __future__ import annotations
from typing import Any, Dict, List, Optional
from dataclasses import dataclass, field
from pathlib import Path
import zipfile
import os

from app.domain.presentation.models import PresentationDocument, Slide, Component, ComponentType
from app.services.design.theme import Theme
from app.core.logging import get_logger

logger = get_logger(__name__)


@dataclass
class QAResultSummary:
    overall_score: int
    scores: Dict[str, int] = field(default_factory=dict)
    passed: bool = False
    checks: List[Dict[str, Any]] = field(default_factory=list)
    issues: List[Dict[str, Any]] = field(default_factory=list)


class QARunner:
    WEIGHTS = {
        "content": 20,
        "educational": 15,
        "visual_hierarchy": 20,
        "readability": 15,
        "design_consistency": 15,
        "editability": 10,
        "technical_validity": 5,
    }

    def __init__(self, theme: Optional[Theme] = None, min_score: int = 85):
        self.theme = theme
        self.min_score = min_score

    def run(self, doc: PresentationDocument, pptx_path: str | None = None) -> QAResultSummary:
        scores: Dict[str, int] = {}
        issues: List[Dict[str, Any]] = []
        checks: List[Dict[str, Any]] = []

        scores["content"] = self._score_content(doc, issues, checks)
        scores["educational"] = self._score_educational(doc, issues, checks)
        scores["visual_hierarchy"] = self._score_visual(doc, issues, checks)
        scores["readability"] = self._score_readability(doc, issues, checks)
        scores["design_consistency"] = self._score_consistency(doc, issues, checks)
        scores["editability"] = self._score_editability(doc, pptx_path, issues, checks)
        scores["technical_validity"] = self._score_technical(doc, pptx_path, issues, checks)

        overall = sum(int(scores[k]) * (self.WEIGHTS[k] / 100) for k in self.WEIGHTS)
        overall_int = int(round(overall))
        passed = overall_int >= self.min_score

        return QAResultSummary(
            overall_score=overall_int,
            scores=scores,
            passed=passed,
            checks=checks,
            issues=issues,
        )

    def _check(self, name: str, ok: bool, weight: int = 1, note: str = "") -> int:
        if ok:
            return weight
        return 0

    def _score_content(self, doc: PresentationDocument, issues, checks) -> int:
        score = 0
        max_score = 100
        if doc.title and len(doc.title) > 5:
            score += 15
        else:
            issues.append({"id": "C01", "severity": "high", "message": "Presentation title is missing or too short"})
        if doc.slide_count() >= 8 and doc.slide_count() <= 20:
            score += 25
        elif doc.slide_count() > 0:
            score += 10
            issues.append({"id": "C02", "severity": "medium",
                           "message": f"Slide count {doc.slide_count()} outside ideal range 8-20"})
        else:
            issues.append({"id": "C03", "severity": "critical", "message": "No slides in presentation"})
        per_slide_ok = 0
        densities = []
        for i, s in enumerate(doc.slides, 1):
            info = s.info_density()
            densities.append(info["density"])
            has_heading = any(c.type in (ComponentType.HEADING, ComponentType.SUBTITLE) for c in s.components)
            if s.main_message or has_heading:
                per_slide_ok += 1
        total_slides = max(1, doc.slide_count())
        score += int(35 * per_slide_ok / total_slides)
        overloaded = densities.count("OVERLOADED")
        if overloaded:
            issues.append({"id": "C04", "severity": "high",
                           "message": f"{overloaded} slides are OVERLOADED (information density)"})
        underloaded = densities.count("UNDERLOADED")
        if underloaded > total_slides // 2:
            issues.append({"id": "C05", "severity": "low",
                           "message": f"{underloaded} slides are UNDERLOADED"})
        else:
            score += 25 - min(25, 5 * max(overloaded, underloaded))

        ref_slides = sum(1 for s in doc.slides if s.source_references)
        if ref_slides >= max(1, total_slides // 4):
            score += 0
        score = min(100, score)
        checks.append({"domain": "content", "slides": doc.slide_count(), "score": score})
        return max(0, min(100, score))

    def _score_educational(self, doc: PresentationDocument, issues, checks) -> int:
        score = 0
        purposes = []
        for s in doc.slides:
            p = s.purpose.value if hasattr(s.purpose, "value") else str(s.purpose)
            purposes.append(p)
        has_objectives = any("OBJECTIVE" in p.upper() for p in purposes)
        has_summary = any("SUMMARY" in p.upper() or "TAKEAWAY" in p.upper() for p in purposes)
        has_intro = any("INTRO" in p.upper() or "AGENDA" in p.upper() for p in purposes)
        score += 30 if has_intro else 10
        score += 30 if has_objectives else 0
        score += 30 if has_summary else 5
        total = max(1, doc.slide_count())
        notes_ok = sum(1 for s in doc.slides if s.speaker_notes)
        if notes_ok >= max(1, total // 2):
            score += 10
        score = min(100, score)
        checks.append({"domain": "educational", "objectives": has_objectives, "score": score})
        if not has_objectives:
            issues.append({"id": "E01", "severity": "medium", "message": "No learning objectives slide"})
        if not has_summary:
            issues.append({"id": "E02", "severity": "medium", "message": "No summary or takeaway slide"})
        return max(0, min(100, score))

    def _score_visual(self, doc: PresentationDocument, issues, checks) -> int:
        score = 0
        total = max(1, doc.slide_count())
        layout_types = set()
        for s in doc.slides:
            lt = s.layout.type.value if hasattr(s.layout.type, "value") else str(s.layout.type)
            layout_types.add(lt)
        score += 25 if len(layout_types) >= 4 else 10
        if len(layout_types) <= 2 and total >= 6:
            issues.append({"id": "V01", "severity": "medium", "message": "Low layout diversity"})
        for s in doc.slides:
            has_card_metric = any(c.type in (ComponentType.CARD, ComponentType.METRIC) for c in s.components)
            has_heading = any(c.type == ComponentType.HEADING for c in s.components)
            if has_heading:
                score += int(35 / total)
            if has_card_metric:
                score += int(25 / total)
        component_counts = [len(s.components) for s in doc.slides]
        mean_c = sum(component_counts) / total if total else 0
        if 2 <= mean_c <= 8:
            score += 15
        score = min(100, score)
        checks.append({"domain": "visual", "layout_types": len(layout_types), "score": score})
        return max(0, min(100, score))

    def _score_readability(self, doc: PresentationDocument, issues, checks) -> int:
        score = 100
        problem_words = 0
        total_words = 0
        for s in doc.slides:
            info = s.info_density()
            total_words += info["word_count"]
            if info["density"] == "OVERLOADED":
                score -= 12
                problem_words += info["word_count"]
            elif info["density"] == "DENSE":
                score -= 4
        avg = total_words / max(1, doc.slide_count())
        if avg > 120:
            score -= 10
        if any(s.info_density()["reading_time_minutes"] > 3 for s in doc.slides):
            issues.append({"id": "R01", "severity": "high",
                           "message": "Some slides take more than 3 minutes to read"})
        score = max(0, min(100, score))
        checks.append({"domain": "readability", "avg_words_per_slide": int(avg), "score": score})
        return max(0, min(100, score))

    def _score_consistency(self, doc: PresentationDocument, issues, checks) -> int:
        score = 100
        if not self.theme:
            checks.append({"domain": "consistency", "score": score})
            return max(0, min(100, score))
        colors = set()
        fonts = set()
        for s in doc.slides:
            for c in s.components:
                if c.style.text and c.style.text.font_family:
                    fonts.add(c.style.text.font_family)
        if len(fonts) > 3:
            score -= 15
            issues.append({"id": "CO01", "severity": "low", "message": "More than 3 font families in use"})
        token_colors = self.theme.color_system.to_palette()
        if token_colors:
            score += 0
        score = max(0, min(100, score))
        checks.append({"domain": "consistency", "font_count": len(fonts), "score": score})
        return score

    def _score_editability(self, doc: PresentationDocument, pptx_path: str | None, issues, checks) -> int:
        score = 70
        try:
            if pptx_path and Path(pptx_path).exists():
                from pptx import Presentation as PptxP
                prs = PptxP(pptx_path)
                all_texts = 0
                editable_shapes = 0
                for slide in prs.slides:
                    for sh in slide.shapes:
                        editable_shapes += 1
                        if sh.has_text_frame and sh.text_frame.text.strip():
                            all_texts += 1
                score = 85 if editable_shapes > 5 else 50
                if all_texts >= 3:
                    score += 10
                if self._pptx_is_valid_zip(pptx_path):
                    score += 5
        except Exception as exc:
            issues.append({"id": "ED01", "severity": "medium",
                           "message": f"Editability probe failed: {exc}"})
        score = max(0, min(100, score))
        checks.append({"domain": "editability", "score": score})
        return score

    def _score_technical(self, doc: PresentationDocument, pptx_path: str | None, issues, checks) -> int:
        score = 85
        if pptx_path:
            path = Path(pptx_path)
            if not path.exists():
                issues.append({"id": "T01", "severity": "critical", "message": "PPTX output path does not exist"})
                score = 0
            else:
                size = path.stat().st_size
                if size < 10_000:
                    score = 40
                    issues.append({"id": "T02", "severity": "medium", "message": "PPTX file suspiciously small"})
                else:
                    if self._pptx_is_valid_zip(str(path)):
                        score += 15
                    else:
                        score = 10
                        issues.append({"id": "T03", "severity": "critical",
                                       "message": "PPTX is not a valid zip package"})
        valid_comps = 0
        total_comps = 0
        for s in doc.slides:
            for c in s.components:
                total_comps += 1
                if c.size.width >= 0 and c.size.height >= 0 and not c.validate():
                    valid_comps += 1
        if total_comps and valid_comps == total_comps:
            score = min(100, score + 0)
        checks.append({"domain": "technical", "valid_components": valid_comps, "score": min(100, score)})
        return max(0, min(100, score))

    def _pptx_is_valid_zip(self, path: str) -> bool:
        try:
            with zipfile.ZipFile(path, "r") as z:
                names = z.namelist()
                if not names:
                    return False
                expected = ["[Content_Types].xml", "ppt/presentation.xml"]
                return all(any(n.endswith(e) or n == e for n in names) for e in expected)
        except Exception as exc:
            logger.warning("pptx zip check failed: %s", exc)
            return False
