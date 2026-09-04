from __future__ import annotations
from typing import TYPE_CHECKING
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

if TYPE_CHECKING:
    from pptx.parts.slide import SlidePart
    from pptx.shapes.base import BaseShape


def inches(n: float) -> int:
    return Inches(n)


def points(n: float) -> int:
    return Pt(n)


def rgb_from_hex(hex_str: str) -> RGBColor:
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_textbox(shape, text: str, size_pt: float = 14, bold: bool = False,
                italic: bool = False, color_hex: str = "#0E1B2D",
                align: str = "left", font_family: str = "Calibri",
                line_spacing: float = 1.3) -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(3)
    tf.margin_right = Pt(3)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                 "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}.get(align, PP_ALIGN.LEFT)
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alignment
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line or " "
        run.font.name = font_family
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb_from_hex(color_hex)


def set_multibullet(shape, items, size_pt: float = 14,
                    color_hex: str = "#1F2A44", font_family: str = "Calibri",
                    line_spacing: float = 1.35, spacing_before: float = 4,
                    bullet_color_hex: str | None = None,
                    align: str = "left") -> None:
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                 "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
    bullet_color_hex = bullet_color_hex or color_hex

    def add_items(target_items, paragraph_ref, level=0):
        first = True
        for item in target_items:
            p = paragraph_ref if first else tf.add_paragraph()
            first = False
            p.alignment = alignment
            p.level = min(level, 8)
            p.line_spacing = line_spacing
            p.space_before = Pt(spacing_before) if level == 0 else Pt(2)
            run = p.add_run()
            run.text = f"•  {item.text}"
            run.font.name = font_family
            run.font.size = Pt(max(9, size_pt - level * 1))
            run.font.color.rgb = rgb_from_hex(color_hex)
            if item.bold:
                run.font.bold = True
            if item.sub_items:
                _ = tf.add_paragraph()
                add_items(item.sub_items, None, level + 1)

    if items:
        add_items(items, tf.paragraphs[0])


def add_rectangle(slide, x_in: float, y_in: float, w_in: float, h_in: float,
                  fill_hex: str = "#FFFFFF", line_hex: str | None = None,
                  line_width_pt: float = 0.75, radius: float = 0.08) -> "BaseShape":
    if radius and radius > 0:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                       inches(x_in), inches(y_in),
                                       inches(w_in), inches(h_in))
        try:
            shape.adjustments[0] = min(0.5, radius / min(w_in, h_in) * 4.5 if min(w_in, h_in) else 0.1)
        except Exception:
            pass
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       inches(x_in), inches(y_in),
                                       inches(w_in), inches(h_in))
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = rgb_from_hex(fill_hex)
    line = shape.line
    if line_hex:
        line.color.rgb = rgb_from_hex(line_hex)
        line.width = Pt(line_width_pt)
    else:
        line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_textbox(slide, x_in: float, y_in: float, w_in: float, h_in: float) -> "BaseShape":
    return slide.shapes.add_textbox(inches(x_in), inches(y_in), inches(w_in), inches(h_in))


def set_slide_background(slide, color_hex: str) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb_from_hex(color_hex)


def add_shape_band(slide, y_in: float, h_in: float, color_hex: str,
                   width_in: float = 13.333) -> "BaseShape":
    return add_rectangle(slide, 0, y_in, width_in, h_in, fill_hex=color_hex, line_hex=None, radius=0.0)


def set_speaker_notes(slide, notes: list[str]) -> None:
    if not notes:
        return
    text = "\n\n".join(notes)
    try:
        notes_tf = slide.notes_slide.notes_text_frame
        notes_tf.text = text
    except Exception:
        pass
